"""Tabular heat-exchanger datasheet and PDF/XLSX exporters.

Export values are sourced from the result payload and its input specifications.
Unavailable fields are intentionally left blank; this module does not estimate
construction details that the design calculation did not produce.
"""
from __future__ import annotations

from typing import Any, Dict, Iterable, List, Tuple


# Datasheet rows follow the user's reference sheet (rows 10-57).
PERFORMANCE_FIELDS: List[Tuple[int, str, str, str]] = [
    (10, "Fluid allocation", "", "fluid_allocation"),
    (11, "Fluid name", "", "fluid_name"),
    (12, "Fluid quantity, Total", "kg/h", "mass_flow"),
    (13, "Vapor (In/Out)", "kg/h", "vapor_flow"),
    (14, "Liquid", "kg/h", "liquid_flow"),
    (15, "Noncondensable", "kg/h", "noncondensable_flow"),
    (17, "Temperature (In/Out)", "°C", "temperature"),
    (18, "Dew / Bubble point", "°C", "dew_bubble_point"),
    (19, "Density", "kg/m³", "density"),
    (20, "Viscosity", "cP", "viscosity"),
    (21, "Molecular wt, Vap", "", "molecular_weight_vapor"),
    (22, "Molecular wt, NC", "", "molecular_weight_noncondensable"),
    (23, "Specific heat", "kJ/(kg°C)", "specific_heat"),
    (24, "Thermal conductivity", "W/(mK)", "thermal_conductivity"),
    (25, "Latent heat", "kJ/kg", "latent_heat"),
    (26, "Pressure", "mmH2O(g)", "pressure"),
    (27, "Velocity", "m/s", "velocity"),
    (28, "Pressure drop, allow./calc.", "mmH2O", "pressure_drop"),
    (29, "Fouling resist. (min)", "m²K/W", "fouling_resistance"),
    (30, "Heat exchanged", "kcal/h", "heat_duty"),
    (31, "Transfer rate, Service", "W/(m²K)", "heat_transfer_coefficient"),
]
CONSTRUCTION_FIELDS: List[Tuple[int, str, str]] = [
    (34, "Design pressure", "kgf/cm²"),
    (35, "Design temperature", "°C"),
    (36, "Number passes per shell", ""),
    (37, "Corrosion allowance", "mm"),
    (38, "Connections", ""),
    (39, "Size/rating", ""),
    (40, "Tube dimensions (inlet/outlet/intermediate)", "mm"),
    (41, "Tube No. / OD / average thickness / length / pitch", "mm"),
    (42, "Tube type / material / tube pattern", ""),
    (43, "Shell ID / OD / thickness", "mm"),
    (44, "Channel or bonnet / shell cover / channel cover", ""),
    (45, "Tubesheet (stationary/floating)", ""),
    (46, "Floating head cover / impingement protection", ""),
    (47, "Baffle crossing / type / cut / spacing", "mm"),
    (48, "Baffle-long / seal type / inlet", "mm"),
    (49, "Supports-tube / U-bend / type", ""),
    (50, "Bypass seal / tube-tubesheet joint / groove-expand", ""),
    (51, "Expansion joint / type", ""),
    (52, "RhoV² / inlet nozzle / bundle entrance / bundle exit", "kg/(m·s²)"),
    (53, "Gaskets - shell side / tube side", ""),
    (54, "Floating head", ""),
    (55, "Code requirements / TEMA class", ""),
    (56, "Weight (filled with water / bundle)", "kg"),
    (57, "Remarks", ""),
]


def _plain(value: Any) -> Any:
    if value is None:
        return None
    if hasattr(value, "value") and not isinstance(value, (str, bytes)):
        try:
            return value.value
        except Exception:
            pass
    if isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, dict):
        return "; ".join(f"{k}: {_plain(v)}" for k, v in value.items())
    if isinstance(value, (list, tuple)):
        return ", ".join(str(_plain(v)) for v in value)
    return str(value)


def _first(data: Dict[str, Any], *keys: str) -> Any:
    for key in keys:
        value = data.get(key)
        if value is not None:
            return value
    return None


def _stream_data(data: Dict[str, Any], side: str) -> Dict[str, Any]:
    """Normalize MaterialStream inputs and optional precomputed properties."""
    streams = data.get("stream_properties") or data.get("streams") or {}
    candidate = streams.get(side, {}) if isinstance(streams, dict) else {}
    if isinstance(candidate, dict) and candidate:
        return candidate
    stream = data.get("hot_in" if side == "hot" else "cold_in")
    if stream is None:
        return {}
    component = getattr(stream, "component", None)
    name = getattr(component, "name", None) or getattr(stream, "name", None)
    props: Dict[str, Any] = {"name": name}
    def read_unit(value: Any, unit: str) -> Any:
        if value is None:
            return None
        try:
            converted = value.to(unit) if hasattr(value, "to") else value
            return getattr(converted, "original_value", getattr(converted, "value", converted))
        except Exception:
            return value
    mass_flow = getattr(stream, "mass_flow", None)
    if callable(mass_flow):
        try:
            mass_flow = mass_flow()
        except Exception:
            mass_flow = None
    props["m_dot_kg_h"] = read_unit(mass_flow, "kg/h")
    outlet = data.get("hot_out" if side == "hot" else "cold_out")
    t_in = read_unit(getattr(stream, "temperature", None), "C")
    t_out = read_unit(getattr(outlet, "temperature", None), "C") if outlet is not None else None
    props["temperature"] = (t_in, t_out)
    props["pressure"] = read_unit(getattr(stream, "pressure", None), "mmH2O")
    props["density"] = read_unit(getattr(stream, "density", None), "kg/m3")
    cp = getattr(stream, "specific_heat", None)
    props["specific_heat"] = read_unit(cp, "kJ/kgK")
    props["molecular_weight_vapor"] = getattr(component, "molecular_weight", None)
    for attr, keys in (
        ("viscosity", ("viscosity",)),
        ("thermal_conductivity", ("thermal_conductivity", "conductivity")),
    ):
        value = None
        for key in keys:
            raw = getattr(component, key, None) if component is not None else None
            if callable(raw):
                try:
                    raw = raw()
                except Exception:
                    raw = None
            if raw is not None:
                value = raw
                break
        if attr == "viscosity":
            props[attr] = read_unit(value, "cP")
        elif attr == "thermal_conductivity":
            props[attr] = read_unit(value, "W/mK")
        else:
            props[attr] = value
    return props


def build_datasheet(data: Dict[str, Any]) -> List[Dict[str, Any]]:
    """Return ordered rows with row number, label, unit, shell-side and tube-side."""
    specs = data.get("specs") or {}
    geometry = data.get("geometry") or {}
    hot = _stream_data(data, "hot")
    cold = _stream_data(data, "cold")
    tube_side = data.get("tube_side", "hot")
    # Prefer the exchanger's actual assignment when it is reported.
    tube_fluid = data.get("tube_side_fluid")
    if tube_fluid:
        hot_name = hot.get("name")
        cold_name = cold.get("name")
        if str(tube_fluid) == str(hot_name):
            tube_side = "hot"
        elif str(tube_fluid) == str(cold_name):
            tube_side = "cold"
    sides = {"Shell Side": cold if tube_side == "hot" else hot,
             "Tube Side": hot if tube_side == "hot" else cold}

    def get_side(side: str, field: str) -> Any:
        props = sides[side]
        aliases = {
            "fluid_name": ("name", "fluid_name", "component"),
            "mass_flow": ("m_dot_kg_h", "mass_flow_kg_h", "m_dot", "mass_flow"),
            "temperature": ("temperature_in_out", "temperature", "t_in_out"),
            "density": ("density_kg_m3", "density"),
            "viscosity": ("viscosity_cp", "viscosity"),
            "specific_heat": ("cp_kj_kg_k", "specific_heat", "cp"),
            "thermal_conductivity": ("k_w_m_k", "thermal_conductivity", "k"),
            "pressure": ("pressure", "pressure_mmh2o"),
            "velocity": ("velocity",),
            "pressure_drop": ("pressure_drop", "dp"),
            "latent_heat": ("latent_heat",),
        }
        return _first(props, *aliases.get(field, (field,)))

    rows: List[Dict[str, Any]] = []
    for number, label, unit, field in PERFORMANCE_FIELDS:
        if field == "fluid_allocation":
            shell = "Cold" if tube_side == "hot" else "Hot"
            tube = "Hot" if tube_side == "hot" else "Cold"
        elif field == "heat_duty":
            q = _first(data, "Q", "heat_duty", "q_watts")
            try:
                q = float(q.to("W").value if hasattr(q, "to") else q)
                q = q * 0.8598452279
            except (TypeError, ValueError, AttributeError):
                q = None
            shell = tube = q
        elif field == "heat_transfer_coefficient":
            shell = tube = _first(data, "U_calculated", "U", "overall_U")
        elif field == "fouling_resistance":
            shell = _first(specs, "shell_fouling_resistance", "shell_fouling")
            tube = _first(specs, "tube_fouling_resistance", "tube_fouling")
        elif field in {"velocity", "pressure_drop"}:
            shell_value = data.get("shell_velocity" if field == "velocity" else "shell_dp")
            tube_value = data.get("tube_velocity" if field == "velocity" else "tube_dp")
            if field == "pressure_drop":
                shell = (data.get("shell_dp_limit", specs.get("shell_dp")), shell_value)
                tube = (data.get("tube_dp_limit", specs.get("tube_dp")), tube_value)
            else:
                shell, tube = shell_value, tube_value
        else:
            shell, tube = get_side("Shell Side", field), get_side("Tube Side", field)
        rows.append({"row": number, "label": label, "unit": unit,
                     "shell_side": _plain(shell), "tube_side": _plain(tube)})

    # Construction rows are intentionally input/result-backed; unknown values stay blank.
    construction_keys = {
        34: ("design_pressure", "design_pressure"),
        35: ("design_temperature", "design_temperature"),
        36: ("shell_passes", "shell_passes"),
        37: ("corrosion_allowance", "corrosion_allowance"),
        38: ("connections", "connections"),
        39: ("size_rating", "size_rating"),
        40: ("tube_dimensions", "tube_dimensions"),
        41: ("tube_count", "tube_count"),
        42: ("tube_type", "tube_type"),
        43: ("shell_diameter", "shell_diameter"),
        44: ("channel_or_bonnet", "channel_or_bonnet"),
        45: ("tubesheet", "tubesheet"),
        46: ("floating_head_cover", "floating_head_cover"),
        47: ("baffle_cut", "baffle_cut"),
        48: ("baffle_spacing", "baffle_spacing"),
        49: ("supports", "supports"),
        50: ("tube_tubesheet_joint", "tube_tubesheet_joint"),
        51: ("expansion_joint", "expansion_joint"),
        52: ("rho_v2", "rho_v2"),
        53: ("gaskets", "gaskets"),
        54: ("floating_head", "floating_head"),
        55: ("code_requirements", "tema_class"),
        56: ("weight_filled", "weight_bundle"),
        57: ("remarks", "remarks"),
    }
    for number, label, unit in CONSTRUCTION_FIELDS:
        key_a, key_b = construction_keys[number]
        value_a = geometry.get(key_a)
        if value_a is None:
            value_a = specs.get(key_a)
        if value_a is None:
            value_a = data.get(key_a)
        value_b = geometry.get(key_b)
        if value_b is None:
            value_b = specs.get(key_b)
        if value_b is None:
            value_b = data.get(key_b)
        # For common geometry values use actual result keys when present.
        if number == 36:
            value_a = _first(data, "shell_passes", specs.get("shell_passes"))
            value_b = _first(data, "tube_passes", specs.get("tube_passes"))
        elif number == 41:
            value_a = _first(data, "tube_count", specs.get("tube_count"))
            value_b = _first(data, "tube_length", specs.get("tube_length"))
        elif number == 43:
            value_a = _first(data, "shell_diameter", specs.get("shell_diameter"))
            value_b = _first(data, "shell_thickness", specs.get("shell_thickness"))
        elif number == 47:
            value_a = _first(data, "baffle_cut", specs.get("baffle_cut"))
            value_b = _first(data, "baffle_spacing", specs.get("baffle_spacing"))
        rows.append({"row": number, "label": label, "unit": unit,
                     "shell_side": _plain(value_a), "tube_side": _plain(value_b)})
    return rows


def _excel_value(value: Any) -> Any:
    value = _plain(value)
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def export_excel(data: Dict[str, Any], path: str) -> str:
    """Write an editable Excel datasheet. Requires optional dependency openpyxl."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, Side, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise ImportError("Excel export requires openpyxl. Install it with 'pip install openpyxl'.") from exc

    wb = Workbook()
    ws = wb.active
    ws.title = "HX Datasheet"
    metadata = data.get("datasheet_metadata") or {}
    ws.merge_cells("A1:E1")
    ws["A1"] = metadata.get("title", "HEAT EXCHANGER SPECIFICATION SHEET")
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")
    meta_rows = [
        ("Company", metadata.get("company"), "Location", metadata.get("location"), "Service", metadata.get("service_of_unit")),
        ("Item No.", metadata.get("item_no"), "Prepared by", metadata.get("prepared_by"), "Date", metadata.get("date")),
        ("Rev No.", metadata.get("revision"), "Job No.", metadata.get("job_no"), "Size", metadata.get("size")),
        ("Type", metadata.get("type"), "Connected in", metadata.get("connected_in"), "Arrangement", metadata.get("arrangement")),
        ("Surface", metadata.get("surface"), "Shells/unit", metadata.get("shells_per_unit"), "Shells", metadata.get("shells")),
    ]
    for row_no, fields in enumerate(meta_rows, 2):
        ws.cell(row_no, 1, fields[0]); ws.cell(row_no, 2, fields[1])
        ws.cell(row_no, 3, fields[2]); ws.cell(row_no, 4, fields[3])
        ws.cell(row_no, 5, f"{fields[4]}: {fields[5] if fields[5] is not None else ''}")
    ws.merge_cells("A7:E7")
    ws["A7"] = "PERFORMANCE OF ONE UNIT"
    ws["A7"].font = Font(bold=True)
    ws["A7"].alignment = Alignment(horizontal="center")
    headers = ["Ref. row", "Description", "Unit", "Shell Side", "Tube Side"]
    for col, value in enumerate(headers, 1):
        cell = ws.cell(8, col, value)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    rows = build_datasheet(data)
    current = 9
    sections = {34: "CONSTRUCTION OF ONE SHELL"}
    thin = Side(style="thin", color="333333")
    for row in rows:
        if row["row"] in sections:
            ws.merge_cells(start_row=current, start_column=1, end_row=current, end_column=5)
            cell = ws.cell(current, 1, sections[row["row"]])
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
            current += 1
        vals = [row["row"], row["label"], row["unit"], _excel_value(row["shell_side"]), _excel_value(row["tube_side"])]
        for col, value in enumerate(vals, 1):
            cell = ws.cell(current, col, value)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        current += 1
    for col, width in enumerate([10, 45, 18, 25, 25], 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "D5"
    ws.sheet_view.showGridLines = False
    wb.save(path)
    return path


def export_pdf(data: Dict[str, Any], path: str) -> str:
    """Write a landscape PDF datasheet. Requires optional dependency reportlab."""
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    except ImportError as exc:
        raise ImportError("PDF export requires reportlab. Install it with 'pip install reportlab'.") from exc

    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(path, pagesize=landscape(A4), rightMargin=22, leftMargin=22, topMargin=22, bottomMargin=22)
    story = [Paragraph("PROCESSPI — HEAT EXCHANGER DATASHEET", styles["Title"]),
             Paragraph(f"Type: {data.get('hx_type', 'Heat exchanger')} &nbsp;&nbsp; Status: {data.get('status', '')}", styles["Normal"]),
             Spacer(1, 8)]
    table_data = [["Ref.", "Description", "Unit", "Shell Side", "Tube Side"]]
    sections = {10: "PERFORMANCE OF ONE UNIT", 34: "CONSTRUCTION OF ONE SHELL"}
    for row in build_datasheet(data):
        if row["row"] in sections:
            table_data.append([sections[row["row"]], "", "", "", ""])
        table_data.append([str(row["row"]), row["label"], row["unit"],
                           "" if row["shell_side"] is None else str(row["shell_side"]),
                           "" if row["tube_side"] is None else str(row["tube_side"])])
    table = Table(table_data, colWidths=[36, 250, 72, 170, 170], repeatRows=1)
    style = [
        ("GRID", (0, 0), (-1, -1), 0.35, colors.black),
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightblue),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
    ]
    for idx, row in enumerate(table_data):
        if len(row) and row[0] in sections.values():
            style.extend([("SPAN", (0, idx), (-1, idx)),
                          ("BACKGROUND", (0, idx), (-1, idx), colors.lightgrey),
                          ("FONTNAME", (0, idx), (-1, idx), "Helvetica-Bold"),
                          ("ALIGN", (0, idx), (-1, idx), "CENTER")])
    table.setStyle(TableStyle(style))
    story.append(table)
    doc.build(story)
    return path
